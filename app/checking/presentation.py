from pptx import Presentation as PptxPresentation

class Presentation:
    def __init__(self, file_path):
        self._presentation = PptxPresentation(file_path)

    @property
    def slides_titles(self):
        titles = []
        for slide in self._presentation.slides:
            try:
                titles.append(slide.shapes.title.text)
            except:
                titles.append('')
        return titles

